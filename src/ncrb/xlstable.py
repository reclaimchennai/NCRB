"""Read NCRB tables published as Excel workbooks (.xlsx / .xls).

Spreadsheets need no geometry: the grid is given. The work is to find where
each table starts (a sheet can stack several), separate title / header / body
rows, and resolve merged header cells into a header path per column.
"""

from __future__ import annotations

import re
from pathlib import Path

from .assemble import Table, TableRow
from .pdftable import CONTD_RE, NOTE_START_RE, TABLE_RE, is_cell

ANCHOR_CELL_RE = re.compile(r"^[\[(]\s*(\d{1,3})\s*[\])]$")


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if v != v:  # NaN
            return ""
        return str(int(v)) if v == int(v) else repr(round(v, 6))
    return re.sub(r"\s+", " ", str(v)).strip()


def load_grids(path: str) -> list[tuple[str, list[list[str]], list[tuple[int, int, int, int]]]]:
    """(sheet name, rows of cell text, merged ranges as r0, r1, c0, c1 half-open) per sheet."""
    out = []
    if Path(path).suffix.lower() == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(path, data_only=True)
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                continue
            grid = [[_text(c) for c in row] for row in ws.iter_rows(values_only=True)]
            merged = [(r.min_row - 1, r.max_row, r.min_col - 1, r.max_col) for r in ws.merged_cells.ranges]
            out.append((ws.title, grid, merged))
    else:
        import xlrd

        try:
            wb = xlrd.open_workbook(path, formatting_info=True)
        except Exception:
            # some workbooks on the site have a damaged container; the cell data is usually still readable
            wb = xlrd.open_workbook(path, ignore_workbook_corruption=True)
        for sh in wb.sheets():
            if getattr(sh, "visibility", 0):
                continue
            grid = [[_text(v) for v in sh.row_values(i)] for i in range(sh.nrows)]
            out.append((sh.name, grid, list(sh.merged_cells)))
    return out


def _kind(row: list[str]) -> str:
    cells = [(j, v) for j, v in enumerate(row) if v]
    if not cells:
        return "empty"
    anchors = [v for _, v in cells if ANCHOR_CELL_RE.match(v)]
    if len(anchors) >= 2 and len(anchors) >= 0.7 * len(cells):
        return "anchor"
    nums = [v for _, v in cells if is_cell(v)]
    texts = [v for _, v in cells if not is_cell(v)]
    # a title repeated over each block of a wide table ('TABLE 3B.2', 'TABLE 3B.2 (Contd.)') is still one title
    distinct = {re.sub(r"\s+", " ", CONTD_RE.sub(" ", v)).strip(" -–.").lower() for v in texts}
    if len(distinct) <= 1 and not nums:
        return "single"  # title, section heading or note
    if texts and len(nums) >= 3 and all(re.fullmatch(r"(19|20)\d\d(-\d\d)?", v) for v in nums):
        return "header"  # a heading row of years: 'Sl | Crime | 2001 | 2002 | ...'
    if len(nums) >= 2 and len(nums) >= 0.5 * len(cells):
        return "data"  # label cells may repeat when a wide table is laid out in blocks side by side
    if len(nums) >= 1 and len(nums) >= 0.5 * len(cells) - 1 and len(texts) <= 3:
        return "data"
    return "header" if len(texts) >= 2 else "data"


def tables_from_grid(grid: list[list[str]], merged: list[tuple[int, int, int, int]]) -> list[Table]:
    width = max((len(r) for r in grid), default=0)
    grid = [r + [""] * (width - len(r)) for r in grid]
    owner: dict[tuple[int, int], tuple[int, int]] = {}
    for r0, r1, c0, c1 in merged:
        for i in range(r0, r1):
            for j in range(c0, c1):
                owner[(i, j)] = (r0, c0)

    def val(i: int, j: int) -> str:
        oi, oj = owner.get((i, j), (i, j))
        return grid[oi][oj] if oi < len(grid) and oj < width else ""

    kinds = [_kind(r) for r in grid]
    tables: list[Table] = []
    i, n = 0, len(grid)
    while i < n:
        # --- title: single-cell rows up to the header
        titles: list[str] = []
        while i < n and kinds[i] in ("empty", "single"):
            if kinds[i] == "single":
                titles.append(next(v for v in grid[i] if v))
            i += 1
        if i >= n:
            break
        # --- header rows, optional anchor
        h0 = i
        while i < n:
            if kinds[i] == "header":
                i += 1
            elif kinds[i] == "empty" and next((kinds[m] for m in range(i + 1, min(n, i + 4)) if kinds[m] != "empty"), "") in ("header", "anchor"):
                i += 1  # a blank row inside a tall merged heading
            else:
                break
        header_rows = [r for r in range(h0, i) if kinds[r] == "header"]
        ids: dict[int, int] = {}
        if i < n and kinds[i] == "anchor":
            for j, v in enumerate(grid[i]):
                m = ANCHOR_CELL_RE.match(v)
                if m:
                    ids[j] = int(m.group(1))
            i += 1
        if not ids and i < n and kinds[i] == "data":
            # column numbers written without brackets: a row of consecutive integers under the header
            vals = [(j, v) for j, v in enumerate(grid[i]) if v]
            nums = [int(v) for _, v in vals if re.fullmatch(r"\d{1,3}", v)]
            if len(nums) == len(vals) >= 3 and nums[0] <= 3 and all(b - a == 1 for a, b in zip(nums, nums[1:])):
                ids = {j: int(v) for j, v in vals}
                i += 1
        if not header_rows and not ids:
            # data with no heading: attach to the previous table if it has the same shape
            if not tables:
                i += 1
                continue
        # --- body
        b0 = i
        while i < n:
            k = kinds[i]
            if k == "data":
                i += 1
            elif k == "single":
                # a section heading inside the body, unless it starts the next table or the notes
                text = next(v for v in grid[i] if v)
                nxt = next((kinds[m] for m in range(i + 1, n) if kinds[m] != "empty"), "end")
                if TABLE_RE.match(text) or NOTE_START_RE.match(text) or nxt in ("header", "anchor", "end", "single") and len(text) > 40:
                    break
                if nxt != "data":
                    break
                i += 1
            elif k == "empty" and next((kinds[m] for m in range(i + 1, min(n, i + 4)) if kinds[m] != "empty"), "") == "data":
                i += 1  # blank rows inside the table
            else:
                break
        body_rows = [r for r in range(b0, i) if kinds[r] != "empty"]
        notes: list[str] = []
        singles: list[int] = []
        j = i
        while j < n and kinds[j] in ("empty", "single"):
            if kinds[j] == "single":
                singles.append(j)
            j += 1
        if j < n:
            # single-cell rows directly above the next header are that table's title, not our notes
            while singles and not NOTE_START_RE.match(next(v for v in grid[singles[-1]] if v)):
                j = singles.pop()
        notes = [next(v for v in grid[r] if v) for r in singles]
        i = j
        data_rows = [r for r in body_rows if kinds[r] == "data"]
        if not data_rows:
            continue

        used = [j for j in range(width) if any(grid[r][j] for r in data_rows)]
        if not ids:
            ids = {j: j + 1 for j in used}
        # label columns: leading columns that are text (or a serial number followed by text)
        label_cols: list[int] = []
        for j in used:
            vals = [grid[r][j] for r in data_rows if grid[r][j]]
            numeric = sum(is_cell(v) for v in vals) >= 0.6 * len(vals)
            hdr = " ".join(val(r, j) for r in header_rows).lower()
            serial = bool(re.search(r"\b(s[lr]?\.?\s*no|sl\.?|s\.?\s*n\.?|serial)\b", hdr)) or (
                numeric and not label_cols and all(re.fullmatch(r"\d{1,3}", v) for v in vals) and len(used) > 2
            )
            if not numeric or (serial and not label_cols):
                label_cols.append(j)
            else:
                break
        if not label_cols:
            label_cols = [used[0]]
        data_cols = [j for j in used if j not in label_cols and j in ids]
        # a wide table laid out in blocks repeats its label columns; those copies are not data
        for j in list(data_cols):
            vals = [grid[r][j] for r in data_rows]
            if any(vals == [grid[r][l] for r in data_rows] for l in label_cols):
                data_cols.remove(j)

        title = " ".join(titles)
        m = TABLE_RE.match(title)
        t = Table(re.sub(r"\s+", "", m.group(2)) if m else "", "")
        title = CONTD_RE.sub(" ", title[m.end():] if m else title)
        t.title = re.sub(r"\s+", " ", title).strip(" -–:.")
        for j in data_cols:
            path: list[str] = []
            prev_owner = None
            for r in header_rows:
                o = owner.get((r, j), (r, j))
                v = val(r, j)
                if v and o != prev_owner:
                    path.append(v)
                prev_owner = o
            t.columns[ids[j]] = path
        t.label_header = " / ".join(filter(None, (" ".join(dict.fromkeys(val(r, j) for r in header_rows if val(r, j))) for j in label_cols)))
        serial_col = label_cols[0] if len(label_cols) > 1 else None
        section = ""
        for r in body_rows:
            if kinds[r] == "single":
                section = next(v for v in grid[r] if v)
                continue
            serial = grid[r][serial_col] if serial_col is not None and re.fullmatch(r"\d{1,4}[.)]?|\(?[ivxlc]+\)|[A-Za-z][.)]?", grid[r][serial_col]) else ""
            label = " ".join(grid[r][j] for j in label_cols if grid[r][j] and not (j == serial_col and serial))
            cells = {ids[j]: grid[r][j] for j in data_cols if grid[r][j]}
            t.rows.append(TableRow(section, serial.rstrip(".)"), label, cells, 0))
        t.notes = notes
        if not (t.columns and t.rows):
            continue
        prev = tables[-1] if tables else None
        if (
            prev is not None
            and not t.title
            and not t.table_no
            and not set(t.columns) & set(prev.columns)
            and [r.label for r in t.rows] == [r.label for r in prev.rows]
        ):
            # same rows, further columns: the table was laid out in blocks down the sheet
            prev.columns.update(t.columns)
            for a, b in zip(prev.rows, t.rows):
                a.cells.update(b.cells)
            prev.notes += [x for x in t.notes if x not in prev.notes]
        else:
            tables.append(t)
    return tables


def split_side_by_side(grid: list[list[str]], merged: list[tuple[int, int, int, int]]):
    """Cut a sheet into vertical strips when different tables sit next to each other.

    NCRB's exported sheets put 'TABLE 3B.2 (i)' and 'TABLE 3B.2 (ii)' side by
    side, each heading repeated over the blocks of its table. The first row
    tells where one table ends and the next begins.
    """
    top = next((r for r in grid if any(r)), None)
    if top is None:
        return [(grid, merged)]
    starts: list[int] = []
    last = None
    for j, v in enumerate(top):
        if not v:
            continue
        m = TABLE_RE.match(v)
        if not m:
            return [(grid, merged)]
        key = re.sub(r"\s+", "", m.group(0)).lower()
        if key != last:
            starts.append(j)
            last = key
    if len(starts) < 2:
        return [(grid, merged)]
    width = max(len(r) for r in grid)
    bounds = list(zip(starts, starts[1:] + [width]))
    bounds[0] = (0, bounds[0][1])
    out = []
    for a, b in bounds:
        sub = [r[a:b] for r in grid]
        sub_merged = [(r0, r1, max(c0, a) - a, min(c1, b) - a) for r0, r1, c0, c1 in merged if c0 < b and c1 > a]
        out.append((sub, sub_merged))
    return out


def extract_workbook(path: str) -> list[tuple[str, Table]]:
    out = []
    for name, grid, merged in load_grids(path):
        for sub, sub_merged in split_side_by_side(grid, merged):
            for t in tables_from_grid(sub, sub_merged):
                out.append((name, t))
    return out
